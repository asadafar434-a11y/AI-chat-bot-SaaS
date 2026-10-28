"""Тесты на ВЫДУМАННЫХ данных (ИНН, названия и контакты вымышлены).

Структура XML здесь — предположение (см. parser.py), а не копия настоящего контракта.
"""
import io
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from lead_collector import db, export, pipeline  # noqa: E402
from lead_collector.client import build_request, parse_response  # noqa: E402
from lead_collector.models import detect_kind  # noqa: E402
from lead_collector.parser import parse_contract_xml  # noqa: E402

CONTRACT = """<?xml version="1.0" encoding="UTF-8"?>
<export xmlns="http://zakupki.gov.ru/oos/export/1" xmlns:oos="http://zakupki.gov.ru/oos/types/1">
 <contract>
  <regNum>{num}</regNum>
  <signDate>{date}+03:00</signDate>
  <price>{price}</price>
  <customer><fullName>МБОУ Тестовая школа</fullName></customer>
  <contractSubject>Поставка тестовых канцтоваров</contractSubject>
  <suppliers>
   <supplier>{supplier}</supplier>
  </suppliers>
 </contract>
</export>"""

OOO = """<legalEntityRF>
   <fullName>Общество с ограниченной ответственностью "Ромашка"</fullName>
   <INN>7700000001</INN><KPP>770001001</KPP>
   <contactPhone>+7 (495) 000-00-01</contactPhone><contactEMail>Info@Romashka.Test</contactEMail>
  </legalEntityRF>"""
IP = """<individualPersonRF>
   <lastName>Тестов</lastName><firstName>Тест</firstName><middleName>Тестович</middleName>
   <INN>770000000002</INN>
  </individualPersonRF>"""


def make(num="0100000000000000001", date="2026-10-19", price="150000.50", supplier=OOO):
    return CONTRACT.format(num=num, date=date, price=price, supplier=supplier).encode("utf-8")


def make_zip(*xmls: bytes) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for i, x in enumerate(xmls):
            zf.writestr(f"c{i}.xml", x)
    return buf.getvalue()


class ParserTests(unittest.TestCase):
    def test_fields_ooo(self):
        (rec,) = parse_contract_xml(make(), "77")
        self.assertEqual(rec.reg_number, "0100000000000000001")
        self.assertEqual(rec.sign_date, "2026-10-19")
        self.assertEqual(rec.amount, 150000.50)
        self.assertEqual(rec.customer, "МБОУ Тестовая школа")
        self.assertEqual(rec.subject, "Поставка тестовых канцтоваров")
        self.assertEqual((rec.inn, rec.kpp, rec.kind), ("7700000001", "770001001", "ООО"))
        self.assertEqual(rec.phone, "+7 (495) 000-00-01")
        self.assertEqual(rec.email, "info@romashka.test")
        self.assertEqual(rec.customer_region, "77")

    def test_individual_entrepreneur(self):
        (rec,) = parse_contract_xml(make(supplier=IP))
        self.assertEqual(rec.kind, "ИП")
        self.assertEqual(rec.name, "Тестов Тест Тестович")
        self.assertEqual(rec.phone, "")  # чего нет в файле — не выдумываем
        self.assertEqual(rec.email, "")

    def test_kind_detection(self):
        self.assertEqual(detect_kind("123456789012", "Иванов"), "ИП")
        self.assertEqual(detect_kind("1234567890", 'ООО "Х"'), "ООО")
        self.assertEqual(detect_kind("1234567890", "АО Игрек"), "Юрлицо")

    def test_supplier_without_inn_is_skipped(self):
        bad = "<legalEntityRF><fullName>Без ИНН</fullName></legalEntityRF>"
        self.assertEqual(parse_contract_xml(make(supplier=bad)), [])

    def test_amount_comma(self):
        (rec,) = parse_contract_xml(make(price="1 000,25"))
        self.assertEqual(rec.amount, 1000.25)


class StorageTests(unittest.TestCase):
    def setUp(self):
        self.conn = db.connect(":memory:")

    def store(self, *xmls, min_amount=0.0):
        return pipeline.process_archive(self.conn, make_zip(*xmls), "77", min_amount)

    def test_dedupe_by_inn_and_counter(self):
        self.store(make("A1", "2026-10-18"), make("A2", "2026-10-19", price="200"))
        row = self.conn.execute("SELECT * FROM suppliers").fetchall()
        self.assertEqual(len(row), 1)
        self.assertEqual(row[0]["contracts_count"], 2)
        self.assertEqual(row[0]["last_contract_number"], "A2")
        self.assertEqual(row[0]["last_contract_date"], "2026-10-19")
        self.assertEqual(row[0]["last_amount"], 200.0)

    def test_older_contract_does_not_replace_last(self):
        self.store(make("A2", "2026-10-19"))
        self.store(make("A1", "2026-10-01"))
        row = self.conn.execute("SELECT * FROM suppliers").fetchone()
        self.assertEqual(row["last_contract_number"], "A2")
        self.assertEqual(row["contracts_count"], 2)

    def test_rerun_same_period_makes_no_duplicates(self):
        first = self.store(make("A1"), make("B1", supplier=IP))
        again = self.store(make("A1"), make("B1", supplier=IP))
        self.assertEqual(first.new_contracts, 2)
        self.assertEqual(again.new_contracts, 0)
        self.assertEqual(again.repeated_contracts, 2)
        self.assertEqual(db.stats(self.conn)["всего"], 2)
        counts = [r[0] for r in self.conn.execute("SELECT contracts_count FROM suppliers")]
        self.assertEqual(counts, [1, 1])

    def test_status_survives_update(self):
        self.store(make("A1"))
        db.set_contact_status(self.conn, "7700000001", "связались")
        self.store(make("A2", "2026-10-20"))
        status = self.conn.execute("SELECT contact_status FROM suppliers").fetchone()[0]
        self.assertEqual(status, "связались")

    def test_default_status_and_bad_status(self):
        self.store(make("A1"))
        self.assertEqual(self.conn.execute("SELECT contact_status FROM suppliers").fetchone()[0], "не связывались")
        with self.assertRaises(ValueError):
            db.set_contact_status(self.conn, "7700000001", "позвонить")

    def test_min_amount(self):
        rep = self.store(make("A1", price="50"), make("A2", price="500"), min_amount=100)
        self.assertEqual(rep.skipped_by_amount, 1)
        self.assertEqual(db.stats(self.conn)["всего"], 1)

    def test_period_filter(self):
        rep = pipeline.process_archive(self.conn, make_zip(make("A1", "2026-09-01")), "77", 0,
                                       ("2026-10-19", "2026-10-19"))
        self.assertEqual(rep.new_contracts, 0)

    def test_broken_xml_is_counted_not_fatal(self):
        rep = self.store(b"<export><contract>", make("A1"))
        self.assertEqual(rep.broken_files, 1)
        self.assertEqual(rep.new_contracts, 1)

    def test_delete_by_inn_removes_everything(self):
        self.store(make("A1"), make("A2", supplier=IP))
        self.assertTrue(db.delete_supplier(self.conn, "770000000002"))
        self.assertEqual(self.conn.execute("SELECT COUNT(*) FROM suppliers WHERE inn='770000000002'").fetchone()[0], 0)
        self.assertEqual(self.conn.execute("SELECT COUNT(*) FROM contracts WHERE inn='770000000002'").fetchone()[0], 0)
        self.assertEqual(db.stats(self.conn)["всего"], 1)
        self.assertFalse(db.delete_supplier(self.conn, "770000000002"))

    def test_delete_leaves_no_trace_in_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "t.sqlite"
            conn = db.connect(path)
            pipeline.process_archive(conn, make_zip(make("A1", supplier=IP)), "77", 0)
            conn.close()
            conn = db.connect(path)
            db.delete_supplier(conn, "770000000002")
            conn.close()
            raw = path.read_bytes()
            self.assertNotIn("Тестов".encode("utf-8"), raw)
            self.assertNotIn(b"770000000002", raw)


class ExportTests(unittest.TestCase):
    def test_export_files_and_excluded_statuses(self):
        conn = db.connect(":memory:")
        pipeline.process_archive(conn, make_zip(make("A1"), make("B1", supplier=IP)), "77", 0)
        db.set_contact_status(conn, "770000000002", "не беспокоить")
        with tempfile.TemporaryDirectory() as tmp:
            csv_path, xlsx_path = export.export_all(conn, Path(tmp), stamp="t")
            text = csv_path.read_text(encoding="utf-8-sig")
            self.assertIn("7700000001", text)
            self.assertNotIn("770000000002", text)  # «не беспокоить» не выгружаем
            self.assertTrue(xlsx_path.stat().st_size > 0)
            csv_all, _ = export.export_all(conn, Path(tmp), include_all=True, stamp="all")
            self.assertIn("770000000002", csv_all.read_text(encoding="utf-8-sig"))

    def test_repeat_export_only_two_plus_and_chosen_kinds(self):
        conn = db.connect(":memory:")
        pipeline.process_archive(conn, make_zip(
            make("A1"), make("A2", "2026-10-20"),                 # ООО: 2 контракта
            make("B1", supplier=IP),                              # ИП: 1 контракт
            make("C1", supplier=IP.replace("770000000002", "770000000003")),
            make("C2", supplier=IP.replace("770000000002", "770000000003")),  # ИП: 2 контракта
        ), "77", 0)
        with tempfile.TemporaryDirectory() as tmp:
            csv_path, _ = export.export_all(conn, Path(tmp), stamp="t", min_contracts=2,
                                            kinds=("ООО", "ИП"), prefix="repeat")
            text = csv_path.read_text(encoding="utf-8-sig")
        self.assertIn("7700000001", text)
        self.assertIn("770000000003", text)
        self.assertNotIn("770000000002", text)  # один контракт — не повтор
        self.assertEqual(csv_path.name, "repeat_t.csv")

    def test_contacts_export_name_and_email_only(self):
        conn = db.connect(":memory:")
        pipeline.process_archive(conn, make_zip(make("A1"), make("B1", supplier=IP)), "77", 0)
        with tempfile.TemporaryDirectory() as tmp:
            path, count = export.export_contacts(conn, Path(tmp), stamp="t")
            from openpyxl import load_workbook

            rows_ = list(load_workbook(path).active.iter_rows(values_only=True))
        self.assertEqual(rows_[0], ("Название компании", "Почта"))
        self.assertEqual(count, 1)  # у ИП в примере нет почты — в короткий список не попадает
        self.assertEqual(rows_[1][1], "info@romashka.test")

    def test_phone_kept_and_formula_neutralised(self):
        self.assertEqual(export._safe("+7 (495) 000-00-01"), "+7 (495) 000-00-01")
        self.assertEqual(export._safe("=HYPERLINK(1)"), "'=HYPERLINK(1)")


class ClientFormatTests(unittest.TestCase):
    def test_request_order_and_escape(self):
        body = build_request("77", "2026-10-19", "RC", "contract", "id-1")
        self.assertLess(body.index("<orgRegion>"), body.index("<subsystemType>"))
        self.assertLess(body.index("<subsystemType>"), body.index("<documentType44>"))
        self.assertIn("<exactDate>2026-10-19</exactDate>", body)

    def test_response_parsing(self):
        ok = b"<r xmlns='x'><dataInfo><archiveUrl>https://a/1.zip</archiveUrl></dataInfo></r>"
        self.assertEqual(parse_response(ok), (["https://a/1.zip"], False))
        self.assertEqual(parse_response(b"<r><noData>true</noData></r>"), ([], True))


if __name__ == "__main__":
    unittest.main()
