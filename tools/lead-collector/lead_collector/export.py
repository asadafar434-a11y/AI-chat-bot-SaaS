"""Выгрузка в CSV (для Excel: UTF-8 с BOM, разделитель «;») и в Excel (.xlsx)."""
from __future__ import annotations

import csv
import re
import sqlite3
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from .config import EXCLUDED_FROM_EXPORT, EXPORT_DIR
from .models import is_personal_data

COLUMNS = [
    ("inn", "ИНН"), ("kpp", "КПП"), ("name", "Наименование"), ("kind", "Тип"),
    ("pd", "Персональные данные (ИП)"), ("region", "Регион поставщика"),
    ("customer_region", "Регион заказчика"), ("phone", "Телефон"), ("email", "Почта"),
    ("last_contract_date", "Дата последнего контракта"), ("last_contract_number", "Номер контракта"),
    ("last_amount", "Сумма"), ("customer", "Заказчик"), ("subject", "Предмет контракта"),
    ("contracts_count", "Число контрактов"), ("source", "Источник (номер контракта)"),
    ("collected_at", "Дата сбора"), ("contact_status", "Статус контакта"),
]


def _safe(value):
    """Защита от формул в Excel: текст, начинающийся с = + - @, делаем обычным текстом."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        if re.fullmatch(r"\+?[\d\s()\-]+", value):  # телефон вида +7 (495) 123-45-67
            return value
        return "'" + value
    return value


def rows(conn: sqlite3.Connection, include_all: bool = False, min_contracts: int = 1,
         kinds: tuple[str, ...] | None = None) -> list[list]:
    """min_contracts — только поставщики, у которых в базе не меньше стольких контрактов; kinds — только эти типы."""
    where: list[str] = []
    params: list = []
    if not include_all:
        where.append("contact_status NOT IN (%s)" % ",".join("?" * len(EXCLUDED_FROM_EXPORT)))
        params += list(EXCLUDED_FROM_EXPORT)
    if min_contracts > 1:
        where.append("contracts_count >= ?")
        params.append(min_contracts)
    if kinds:
        where.append("kind IN (%s)" % ",".join("?" * len(kinds)))
        params += list(kinds)
    sql = "SELECT * FROM suppliers" + (" WHERE " + " AND ".join(where) if where else "")
    sql += (" ORDER BY contracts_count DESC, last_contract_date DESC, inn" if min_contracts > 1
            else " ORDER BY last_contract_date DESC, inn")
    params = tuple(params)
    out = []
    for r in conn.execute(sql, params):
        d = dict(r)
        d["pd"] = "да" if is_personal_data(d["kind"]) else "нет"
        out.append([_safe(d[key]) if d[key] is not None else "" for key, _ in COLUMNS])
    return out


def export_contacts(conn: sqlite3.Connection, out_dir: Path = EXPORT_DIR, stamp: str | None = None,
                    kinds: tuple[str, ...] | None = None) -> tuple[Path, int]:
    """Короткий список: только «Название компании» и «Почта». Только с почтой, по одной строке на поставщика (ИНН)."""
    out_dir.mkdir(parents=True, exist_ok=True)
    stamp = stamp or date.today().isoformat()
    sql = ("SELECT name, email FROM suppliers WHERE email<>'' AND contact_status NOT IN (%s)"
           % ",".join("?" * len(EXCLUDED_FROM_EXPORT)))
    params: list = list(EXCLUDED_FROM_EXPORT)
    if kinds:
        sql += " AND kind IN (%s)" % ",".join("?" * len(kinds))
        params += list(kinds)
    rows_ = conn.execute(sql + " ORDER BY name", params).fetchall()
    wb = Workbook()
    ws = wb.active
    ws.title = "Контакты"
    ws.append(["Название компании", "Почта"])
    for r in rows_:
        ws.append([_safe(r["name"]), _safe(r["email"])])
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 70
    ws.column_dimensions["B"].width = 40
    path = out_dir / f"contacts_{stamp}.xlsx"
    wb.save(path)
    return path, len(rows_)


def export_all(conn: sqlite3.Connection, out_dir: Path = EXPORT_DIR, include_all: bool = False,
               stamp: str | None = None, min_contracts: int = 1, kinds: tuple[str, ...] | None = None,
               prefix: str = "leads") -> tuple[Path, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    stamp = stamp or date.today().isoformat()
    data = rows(conn, include_all, min_contracts, kinds)
    header = [title for _, title in COLUMNS]

    csv_path = out_dir / f"{prefix}_{stamp}.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh, delimiter=";")
        writer.writerow(header)
        writer.writerows(data)

    xlsx_path = out_dir / f"{prefix}_{stamp}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Поставщики"
    ws.append(header)
    for row in data:
        ws.append(row)
    ws.freeze_panes = "A2"
    for i, (_, title) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(len(title) + 2, 14), 45)
    wb.save(xlsx_path)
    return csv_path, xlsx_path
